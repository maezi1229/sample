"""Step1: 仕入先PDF見積を読み取り、確認用レビューシート(xlsx)を作成する。

Excelアプリは起動せず、openpyxlでレビューシートを直接生成する。
出力されたレビューシートを担当者が元PDFと見比べて「確定値」欄を埋めてから、
次工程（自社見積書テンプレートへの転記）に進む。

使い方:
    python step1_extract_supplier_quote.py <仕入先PDFのパス>
"""
import argparse
import json
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parent.parent.parent))
from quote_automation import config
from quote_automation.extraction.field_parser import parse_draft
from quote_automation.extraction.pdf_reader import read_pdf_text

HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
BOLD = Font(bold=True)


def build_review_workbook(pdf_name: str, draft, method: str) -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "確認シート"

    ws["A1"] = f"仕入先見積 読み取り確認シート（元ファイル: {pdf_name} / 読取方式: {method}）"
    ws["A1"].font = Font(bold=True, size=12)
    ws.merge_cells("A1:D1")

    headers = ["項目", "読取値（機械が読んだ値）", "確定値（元PDFと見比べて正しい値を入力）", "備考"]
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=col, value=text)
        cell.font = BOLD
        cell.fill = HEADER_FILL

    r = 4
    for label, val in [
        ("見積書番号", draft.quote_no),
        ("日付", draft.date),
        ("見積金額（合計）", draft.total_amount),
    ]:
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=2, value=val)
        r += 1

    r += 1
    ws.cell(row=r, column=1, value="【明細】").font = BOLD
    r += 1
    for col, text in enumerate(["品名(候補)", "数量", "単価", "金額"], start=1):
        cell = ws.cell(row=r, column=col, value=text)
        cell.font = BOLD
        cell.fill = HEADER_FILL
    r += 1
    for item in draft.line_items:
        ws.cell(row=r, column=1, value=item.name_guess)
        ws.cell(row=r, column=2, value=item.qty_guess)
        ws.cell(row=r, column=3, value=item.unit_price_guess)
        ws.cell(row=r, column=4, value=item.amount_guess)
        r += 1

    r += 2
    ws.cell(row=r, column=1, value="読み取り生データ（参考・OCR誤読を含む可能性あり）").font = BOLD
    r += 1
    for line in draft.raw_text.splitlines():
        ws.cell(row=r, column=1, value=line)
        r += 1

    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 32
    ws.column_dimensions["D"].width = 20

    return wb


def run(pdf_path: Path) -> Path:
    text, method = read_pdf_text(pdf_path, dpi=config.OCR_DPI, lang=config.TESSERACT_LANG)
    draft = parse_draft(text)
    stem = pdf_path.stem

    (config.OCR_TEXT_DIR / f"{stem}.txt").write_text(text, encoding="utf-8")

    draft_json = {
        "source_pdf": str(pdf_path),
        "method": method,
        "quote_no": draft.quote_no,
        "date": draft.date,
        "total_amount": draft.total_amount,
        "line_items": [item.__dict__ for item in draft.line_items],
    }
    (config.DRAFT_DIR / f"{stem}_draft.json").write_text(
        json.dumps(draft_json, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    wb = build_review_workbook(pdf_path.name, draft, method)
    review_path = config.REVIEW_DIR / f"{stem}_review.xlsx"
    wb.save(review_path)

    print(f"読み取り方式: {method}")
    print(f"見積書番号: {draft.quote_no!r}")
    print(f"日付: {draft.date!r}")
    print(f"見積金額: {draft.total_amount!r}")
    print(f"明細候補: {len(draft.line_items)}件")
    for item in draft.line_items:
        print(f"  - {item.name_guess!r} 数量={item.qty_guess} 単価={item.unit_price_guess} 金額={item.amount_guess}")
    print(f"レビューシート出力先: {review_path}")

    return review_path


def main():
    parser = argparse.ArgumentParser(description="仕入先PDF見積を読み取り、確認用レビューシートを作成する")
    parser.add_argument("pdf_path", type=Path, help="仕入先見積PDFのパス")
    args = parser.parse_args()
    run(args.pdf_path)


if __name__ == "__main__":
    main()
