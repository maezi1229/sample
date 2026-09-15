"""見積り集計表（work/配下のみで管理・Gitにはコミットしない機密ファイル）に
見積り1件分の実績を追記する。

集計表には、日付・売り先（客先）名/担当者・仕入先名/担当者・件名・仕入金額・
売り金額（見積時）・利益率・利益額（見積時）・受注状況・決定金額を記録する。
受注状況（受注/失注/未定）と決定金額（実際に受注した/しなかった金額）は
見積り作成時点ではわからないため、まず「未定」・空欄で記録し、結果が
分かり次第、集計表を直接Excelで開いて更新する運用を想定している
（受注できなかった場合も、今後の分析のために決定金額を記録する）。

使い方（以下はダミーデータの例。実在の客先名・仕入先名・金額は書かない）:
    python update_ledger.py <集計表xlsx> \
        --date 2026/08/14 \
        --customer "サンプル工業株式会社 御中" --contact "総務部　鈴木 様" \
        --supplier "テスト製鋼株式会社" --supplier-contact "田中" \
        --title "ABC123_サンプル部品" \
        --cost-amount 100000 --sell-amount 109000 --profit-rate 9 --profit-amount 9000
"""
import argparse
import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

SHEET_TITLE = "見積り集計"
HEADERS = [
    "年月日", "客先名（売り先）", "客先担当者（売り先）",
    "仕入先名", "仕入先担当者",
    "名称（品名・物件名）",
    "仕入金額", "売り金額（見積時）", "利益率(%)", "利益額（見積時）",
    "受注状況", "決定金額",
]
HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
BOLD = Font(bold=True)
DEFAULT_ORDER_STATUS = "未定"

QUOTE_NUMBER_PREFIX = "HK"


def generate_quote_number(ledger_path: Path, date: datetime.date = None) -> str:
    """見積№を発行する（例:「HK-20260915」）。

    採番ルール: 「HK-」＋見積作成年月日(yyyymmdd)。同日に複数件発行する場合は
    2件目以降に枝番を付ける（2件目は「-2」、3件目は「-3」…）。
    「同日に何件目か」は、見積り集計表（ledger_path）に既に記録されている
    同じ日付の行数を数えて判定する（この関数は新規番号を返すだけで、
    集計表への追記自体は呼び出し側が別途行う）。
    """
    date = date or datetime.date.today()
    date_str = date.strftime("%Y%m%d")
    date_jp = date.strftime("%Y/%m/%d")

    count_today = 0
    ledger_path = Path(ledger_path)
    if ledger_path.exists():
        wb = openpyxl.load_workbook(ledger_path, read_only=True)
        ws = wb[SHEET_TITLE] if SHEET_TITLE in wb.sheetnames else wb.active
        for row in ws.iter_rows(min_row=2, max_col=1, values_only=True):
            if row and row[0] == date_jp:
                count_today += 1

    seq = count_today + 1
    base = f"{QUOTE_NUMBER_PREFIX}-{date_str}"
    return base if seq == 1 else f"{base}-{seq}"


def _ensure_workbook(ledger_path: Path) -> openpyxl.Workbook:
    if ledger_path.exists():
        return openpyxl.load_workbook(ledger_path)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET_TITLE
    for col, text in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=text)
        cell.font = BOLD
        cell.fill = HEADER_FILL
    widths = [12, 28, 20, 24, 16, 34, 14, 14, 10, 14, 12, 14]
    for col, w in zip(range(1, len(HEADERS) + 1), widths):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = w
    return wb


def append_ledger_entry(
    ledger_path: Path,
    *,
    date: str,
    customer_name: str,
    contact_name: str,
    item_title: str,
    cost_amount: float,
    sell_amount: float,
    profit_rate: float,
    profit_amount: float,
    supplier_name: str = "",
    supplier_contact: str = "",
    order_status: str = DEFAULT_ORDER_STATUS,
    decided_amount=None,
) -> Path:
    ledger_path = Path(ledger_path)
    ledger_path.parent.mkdir(parents=True, exist_ok=True)
    wb = _ensure_workbook(ledger_path)
    ws = wb[SHEET_TITLE] if SHEET_TITLE in wb.sheetnames else wb.active

    row = ws.max_row + 1
    values = [
        date, customer_name, contact_name,
        supplier_name, supplier_contact,
        item_title,
        cost_amount, sell_amount, profit_rate, profit_amount,
        order_status, decided_amount,
    ]
    for col, val in enumerate(values, start=1):
        ws.cell(row=row, column=col, value=val)

    wb.save(ledger_path)
    return ledger_path


def main():
    parser = argparse.ArgumentParser(description="見積り集計表に1件分の実績を追記する")
    parser.add_argument("ledger_path", type=Path)
    parser.add_argument("--date", required=True)
    parser.add_argument("--customer", required=True, dest="customer_name")
    parser.add_argument("--contact", required=True, dest="contact_name")
    parser.add_argument("--supplier", default="", dest="supplier_name")
    parser.add_argument("--supplier-contact", default="", dest="supplier_contact")
    parser.add_argument("--title", required=True, dest="item_title")
    parser.add_argument("--cost-amount", required=True, type=float)
    parser.add_argument("--sell-amount", required=True, type=float)
    parser.add_argument("--profit-rate", required=True, type=float)
    parser.add_argument("--profit-amount", required=True, type=float)
    parser.add_argument("--order-status", default=DEFAULT_ORDER_STATUS)
    parser.add_argument("--decided-amount", type=float, default=None)
    args = parser.parse_args()

    result_path = append_ledger_entry(
        args.ledger_path,
        date=args.date,
        customer_name=args.customer_name,
        contact_name=args.contact_name,
        supplier_name=args.supplier_name,
        supplier_contact=args.supplier_contact,
        item_title=args.item_title,
        cost_amount=args.cost_amount,
        sell_amount=args.sell_amount,
        profit_rate=args.profit_rate,
        profit_amount=args.profit_amount,
        order_status=args.order_status,
        decided_amount=args.decided_amount,
    )
    print(f"集計表に追記しました: {result_path}")


if __name__ == "__main__":
    main()
