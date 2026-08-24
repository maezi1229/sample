"""Step3: 利益上乗せ計算が正しいかを検証する。

Step2で転記した見積書(xlsx)をLibreOffice Calc(headless)で実際に再計算させ、
そこで得られた値が「仕入単価 × その行の実効上乗せ率を1000円単位で切り上げた値」と
一致するかをPythonで独立に計算し直して突き合わせる。品目ごとに上乗せ率が
異なりうるため、各行の実効上乗せ率はP列（印刷範囲外の参照用セル、
fill_quote_templateが書き込む）から読み取る。
一致しない場合は、テンプレートの数式が壊れている・値がずれている等の
異常を早期に検知できる。

使い方:
    python step3_verify_calculation.py <Step2で作成したxlsx>
"""
import argparse
import sys
import tempfile
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent.parent))
from quote_automation.template.fill_quote import FREIGHT_ROW, ITEM_ROWS, ROUND_DIGITS, excel_roundup
from quote_automation.template.recalc import recalculate_with_libreoffice


class VerificationError(Exception):
    pass


def run(quote_path: Path) -> bool:
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        recalced_path = recalculate_with_libreoffice(quote_path, tmp / "recalced", tmp / "lo_profile")

        wb = openpyxl.load_workbook(recalced_path, data_only=True)
        ws = wb.active

        default_margin_rate = ws["M17"].value
        print(f"デフォルト利益率設定 (M17): {default_margin_rate}"
              f"（{(default_margin_rate - 1) * 100:.1f}%上乗せ、品目ごとに個別の率が設定されていれば別途表示）")
        print()

        all_ok = True
        total_customer = 0.0
        total_profit = 0.0

        for row in ITEM_ROWS + (FREIGHT_ROW,):
            name = ws[f"C{row}"].value
            qty = ws[f"E{row}"].value
            unit_cost = ws[f"N{row}"].value
            customer_unit_price = ws[f"F{row}"].value
            customer_amount = ws[f"G{row}"].value
            cost_amount = ws[f"O{row}"].value
            profit = ws[f"Q{row}"].value
            item_markup_percent = ws[f"P{row}"].value

            if unit_cost is None:
                continue  # この行は未使用

            if item_markup_percent is None:
                # 客先単価を直接指定した品目（上乗せ率の計算を行っていない）。
                # 独立に「期待される単価」を再現する手段がないため、F列の値を基準に
                # 金額・利益の計算（G=E*F, Q=G-O）が正しいかだけを検証する。
                expected_unit_price = customer_unit_price
                expected_amount = expected_unit_price * qty
                expected_profit = expected_amount - cost_amount
                label = f"{row}行目: {name}（単価直接指定）"
            else:
                expected_unit_price = excel_roundup(unit_cost * (1 + item_markup_percent / 100), ROUND_DIGITS)
                expected_amount = expected_unit_price * qty
                expected_profit = expected_amount - cost_amount
                label = f"{row}行目: {name}（上乗せ率{item_markup_percent}%）"

            ok = (
                customer_unit_price == expected_unit_price
                and customer_amount == expected_amount
                and abs((profit or 0) - expected_profit) < 1e-6
            )
            all_ok = all_ok and ok

            print(f"[{'OK' if ok else 'NG'}] {label}")
            print(f"      仕入単価={unit_cost:,.0f} × 数量{qty} → "
                  f"客先単価={customer_unit_price:,.0f}（期待値={expected_unit_price:,.0f}）")
            print(f"      客先金額={customer_amount:,.0f}（期待値={expected_amount:,.0f}） "
                  f"利益={profit:,.0f}（期待値={expected_profit:,.0f}）")

            total_customer += customer_amount
            total_profit += profit or 0

        grand_total = ws["G26"].value
        grand_profit = ws["Q26"].value
        totals_ok = (
            abs(grand_total - total_customer) < 1e-6
            and abs(grand_profit - total_profit) < 1e-6
        )
        all_ok = all_ok and totals_ok

        print()
        print(f"[{'OK' if totals_ok else 'NG'}] 合計金額: {grand_total:,.0f}円"
              f"（明細合計={total_customer:,.0f}円）")
        print(f"[{'OK' if totals_ok else 'NG'}] 合計利益: {grand_profit:,.0f}円"
              f"（明細合計={total_profit:,.0f}円）")
        print()
        print("検証結果: " + ("すべて一致（OK）" if all_ok else "不一致あり（NG） — テンプレートを確認してください"))

        if not all_ok:
            raise VerificationError("計算結果が期待値と一致しませんでした。")

        return all_ok


def main():
    parser = argparse.ArgumentParser(description="転記済み見積書の利益上乗せ計算を検証する")
    parser.add_argument("quote_path", type=Path)
    args = parser.parse_args()
    run(args.quote_path)


if __name__ == "__main__":
    main()
