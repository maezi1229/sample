"""Step1で出力したレビューシート(xlsx)の「確定値」列を読み戻す。

このモジュールは、人がレビューシートを開いて確認・修正した後の値だけを
信頼するための入口。生のOCR結果を直接Step2に渡すことはしない。
"""
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl


class ReviewNotConfirmedError(Exception):
    """確定値が未入力など、レビューが完了していない場合に送出する。"""


@dataclass
class ConfirmedItem:
    name: str
    qty: float
    unit_price: float


@dataclass
class ConfirmedSupplierQuote:
    quote_no: str
    date: str
    total_amount: float
    items: list = field(default_factory=list)


def _to_number(value) -> float:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).replace(",", "").strip()
    if not s:
        return None
    return float(s)


def load_confirmed_supplier_quote(review_path: Path) -> ConfirmedSupplierQuote:
    wb = openpyxl.load_workbook(review_path, data_only=True)
    ws = wb.active

    labels = {}
    detail_header_row = None
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        a = row[0].value
        if a in ("見積書番号", "日付", "見積金額（合計）"):
            labels[a] = row[2].value  # C列 = 確定値
        if a == "品名(候補)":
            detail_header_row = row[0].row
            break

    missing = [k for k in ("見積書番号", "日付", "見積金額（合計）") if not labels.get(k)]
    if missing:
        raise ReviewNotConfirmedError(
            f"確定値が未入力の項目があります: {', '.join(missing)}。"
            "レビューシートのC列を確認してください。"
        )

    items = []
    if detail_header_row is not None:
        for row in ws.iter_rows(min_row=detail_header_row + 1, max_row=ws.max_row):
            name, qty, unit_price, amount = (c.value for c in row[:4])
            if name is None and qty is None and unit_price is None:
                continue
            if isinstance(name, str) and name.startswith("読み取り生データ"):
                break
            if name is None or unit_price is None:
                continue
            qty_num = _to_number(qty)
            price_num = _to_number(unit_price)
            if qty_num is None:
                raise ReviewNotConfirmedError(
                    f"明細『{name}』の数量が未確定です。レビューシートで数量を入力してください。"
                )
            items.append(ConfirmedItem(name=str(name).strip(), qty=qty_num, unit_price=price_num))

    if not items:
        raise ReviewNotConfirmedError("明細（品名・数量・単価）が1件も確定されていません。")

    return ConfirmedSupplierQuote(
        quote_no=str(labels["見積書番号"]).strip(),
        date=str(labels["日付"]).strip(),
        total_amount=_to_number(labels["見積金額（合計）"]),
        items=items,
    )
