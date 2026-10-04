"""客先向け見積りが受注確定した後、仕入先へ発注する「注文書」を作成する。

見積書（fill_quote.py等）とは逆方向の書類——阪和興業が発注者、仕入先が
受注者になる。仕入先が専用の注文書フォーマットを指定していない場合に
この汎用フォーマットを使う（指定がある場合はそのフォーマットに個別対応する）。

テンプレート: quote_automation/assets/purchase_order_template.xlsx
- 発注者欄（阪和興業側、固定）と発注先欄（仕入先側）を左右に並べたヘッダー
- 品目は4枠（16〜19行目）。1行＝商品名・単価・数量・金額（=単価×数量）
- 消費税・ご注文金額（税込）は手動計算して転記する（仕入先見積りの税込/税別
  表記に合わせる。仕入先側の端数処理ルールが見積書と異なることがあるため、
  自動計算にせず都度確認する）
- 備考欄に、紐づく客先の工事仕様書番号・工事名称があれば記載する
  （無ければ備考欄ごと削除してよい）
- 納品先（客先の現場住所等、仕入先からの直送先）を明記する
- 捺印は見積書と同じ画像・考え方（fill_quote.pyの_add_stamp）を使う
- A4横1ページに収まるよう、印刷設定（landscape・fit to page）を既定で
  テンプレートに設定済み
- 品目が4件を超える場合は、このテンプレートでは収まらない。行を追加で
  挿入すると既存の結合セル・書式が崩れるおそれがあるため、既存の行を
  複製してから書式を手動で揃えるか、一度相談してから対応する
"""
from pathlib import Path

import openpyxl

from quote_automation.template.fill_quote import (
    STAMP_DEFAULT_OFFSET_X_PX,
    STAMP_DEFAULT_SIZE_PX,
    _add_stamp,
)

ITEM_ROWS = (16, 17, 18, 19)

STAMP_DEFAULT_CELL = "J34"
STAMP_DEFAULT_OFFSET_Y_PX = 0

DEFAULT_ORDERER_NAME = "阪和興業株式会社"
DEFAULT_ORDERER_ADDRESS = "〒452-0002　愛知県名古屋市中村区名駅1-1-1　JPタワー名古屋35階"
DEFAULT_ORDERER_TEL = "052-977-3532"
DEFAULT_ORDERER_CONTACT = "前島　拓弥　様"
DEFAULT_SIGNATURE_LINE1 = "阪和興業株式会社"
DEFAULT_SIGNATURE_LINE2 = "名古屋厚板２課　前島　拓弥　印"


class TooManyItemsError(Exception):
    pass


class MissingRequiredFieldError(Exception):
    pass


def _require(value, label: str, missing: list) -> None:
    if value is None or (isinstance(value, str) and not value.strip()):
        missing.append(label)


def fill_purchase_order_sheet(
    ws,
    *,
    supplier_name: str,
    supplier_contact: str,
    supplier_address: str,
    supplier_tel_fax: str,
    supplier_quote_no: str,
    order_date: str,
    order_number: str,
    delivery_date: str,
    items: list,
    tax_amount: float,
    total_with_tax: float,
    delivery_address: str,
    delivery_company: str,
    delivery_contact: str,
    delivery_condition: str = "車上渡し",
    payment_terms: str = "従来通り",
    remarks: str = None,
    stamp_path: Path = None,
    stamp_cell: str = STAMP_DEFAULT_CELL,
) -> None:
    """
    既に開いている1枚のシートに、確定済みの発注データを書き込む（ファイルI/Oなし）。

    items: [{"name": 商品名, "unit_price": 単価, "qty": 数量}, ...] 最大4件まで。
        金額はテンプレート側の固定数式（単価×数量）で自動計算される。
    tax_amount / total_with_tax: 消費税額・税込合計。仕入先見積りの表記
        （税別か税込か、端数処理）に合わせて手動で計算して渡す。
    remarks: 客先の工事仕様書番号など、紐づく資料があれば記載する文言。
        無ければNoneのままにし、備考欄は空欄にする。
    """
    missing = []
    _require(supplier_name, "発注先会社名", missing)
    _require(supplier_contact, "発注先担当者", missing)
    _require(delivery_address, "納品先住所", missing)
    _require(delivery_condition, "納品条件", missing)
    if missing:
        raise MissingRequiredFieldError(
            f"チャットでの確認が必要な項目が未指定です: {', '.join(missing)}。"
        )
    if len(items) > len(ITEM_ROWS):
        raise TooManyItemsError(
            f"テンプレートの品目枠は{len(ITEM_ROWS)}件までです（{len(items)}件指定されました）。"
        )

    ws["B2"] = f"{supplier_name}　御中"
    ws["B3"] = supplier_contact
    ws["H2"] = f"注文日：{order_date}"
    ws["H3"] = f"注文番号：{order_number}"

    ws["C7"] = DEFAULT_ORDERER_NAME
    ws["C8"] = DEFAULT_ORDERER_ADDRESS
    ws["C9"] = DEFAULT_ORDERER_TEL
    ws["C10"] = DEFAULT_ORDERER_CONTACT

    ws["G7"] = supplier_name
    ws["G8"] = supplier_address
    ws["G9"] = supplier_tel_fax
    ws["G10"] = supplier_contact
    ws["G11"] = supplier_quote_no or ""
    ws["D11"] = delivery_date

    ws["D13"] = tax_amount
    ws["H13"] = total_with_tax

    for row, item in zip(ITEM_ROWS, items):
        ws[f"B{row}"] = item["name"]
        ws[f"E{row}"] = item["unit_price"]
        ws[f"G{row}"] = item["qty"]

    if remarks:
        ws["B23"] = remarks

    ws["D26"] = delivery_address
    ws["D27"] = delivery_company
    ws["D28"] = delivery_contact
    ws["D29"] = delivery_condition
    ws["D30"] = payment_terms

    ws["H34"] = DEFAULT_SIGNATURE_LINE1
    ws["H35"] = DEFAULT_SIGNATURE_LINE2

    if stamp_path:
        _add_stamp(
            ws, stamp_path,
            cell=stamp_cell, offset_x_px=STAMP_DEFAULT_OFFSET_X_PX,
            offset_y_px=STAMP_DEFAULT_OFFSET_Y_PX, size_px=STAMP_DEFAULT_SIZE_PX,
        )


def fill_purchase_order_template(template_path: Path, output_path: Path, **kwargs) -> Path:
    """テンプレートファイルを1件分の発注データで埋め、単独のxlsxとして保存する。
    kwargsはfill_purchase_order_sheet()と同じ。"""
    import shutil

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(template_path, output_path)

    wb = openpyxl.load_workbook(output_path)
    fill_purchase_order_sheet(wb.active, **kwargs)

    wb.save(output_path)
    return output_path
