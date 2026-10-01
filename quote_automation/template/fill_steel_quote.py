"""鋼材（板厚×幅×長さの寸法もの）見積書テンプレートへ、確定済みの見積データを転記する。

通常の見積（fill_quote.py、仕入単価×上乗せ率方式）・重量×キロ単価見積
（fill_weight_quote.py）とは別の第三の形式。仕入先（鋼材商社）から
「板厚×幅×長さ×数量」の寸法表形式で原価見積りを受け取る鋼材案件向け
（2026/10/01、前島様の実例・指示をもとにテンプレート化）。

テンプレート: quote_automation/assets/steel_material_quote_template.xlsx
- 明細は8枠（18〜25行目）。1行 = 板厚(B)・幅(C)・長さ(D)・単重量kg(E)・数量(F)・
  総重量kg(G、=E×F)・客先単価円/kg(H)・客先枚単価円(I、=E×H)・客先合計金額円(J、=F×I)
- 印刷範囲（A1:J28）の外側（M〜P列）に原価欄を持つ。仕入先見積りの
  「板厚ごとの枚単価（製品価格）」をO列にそのまま転記するのが入力の起点
  （kg単価は参考値のため、Nは=O/Eの数式でOから逆算する）。
  P16セルに「円/kg乗せ」（客先単価 = 原価単価(円/kg) + この値）を入力すると、
  H〜J列（客先向け表示）とD8/D9（総計金額）まで自動で連動して再計算される。
  ％の上乗せ率（M17のような仕組み）ではなく、円/kgの定額乗せである点が
  通常の見積り（fill_quote.py）と異なる。
- 運賃は既定では客先向け金額（J26）を別途・固定額で入力する
  （乗せ代(P16)とは連動しない。従来通り仕入運賃に固定の掛け率を
  乗せた金額を直接指定する運用。連動させたい場合は都度相談する）。
- 見積№・捺印は通常の見積書と同じルール（fill_quote.pyと同じテンプレート
  由来のレイアウトのため、STAMP_DEFAULT_CELL等もfill_quote.pyと共通）。
"""
from pathlib import Path

import openpyxl

from quote_automation.template.fill_quote import (
    STAMP_DEFAULT_CELL,
    STAMP_DEFAULT_OFFSET_X_PX,
    STAMP_DEFAULT_OFFSET_Y_PX,
    STAMP_DEFAULT_SIZE_PX,
    _add_stamp,
)

ITEM_ROWS = tuple(range(18, 26))  # 8枠、18〜25行目（1行1枠）

QUOTE_NUMBER_CELL = "H4"
QUOTE_NUMBER_LABEL = "見積№"

MARKUP_PER_KG_CELL = "P16"  # 円/kg乗せ（原価単価＋この値＝客先単価）

DEFAULT_PAYMENT_TERMS = "従来通り"
DEFAULT_RECEIVING_PLACE = "貴社車上渡し"


class TooManyItemsError(Exception):
    pass


class MissingRequiredFieldError(Exception):
    pass


def _require(value, label: str, missing: list) -> None:
    if value is None or (isinstance(value, str) and not value.strip()):
        missing.append(label)


def fill_steel_quote_sheet(
    ws,
    *,
    customer_name: str,
    contact_name: str,
    item_title: str,
    items: list,
    markup_per_kg: float,
    freight_customer_price: float = None,
    receiving_place: str = DEFAULT_RECEIVING_PLACE,
    payment_terms: str = DEFAULT_PAYMENT_TERMS,
    note: str = None,
    quote_number: str = None,
    stamp_path: Path = None,
    stamp_cell: str = STAMP_DEFAULT_CELL,
    stamp_offset_x_px: int = STAMP_DEFAULT_OFFSET_X_PX,
    stamp_offset_y_px: int = STAMP_DEFAULT_OFFSET_Y_PX,
    stamp_size_px: int = STAMP_DEFAULT_SIZE_PX,
) -> None:
    """
    既に開いている1枚のシートに、確定済みの見積データを書き込む（ファイルI/Oなし）。

    items: [{"thickness": 板厚, "width": 幅, "length": 長さ, "unit_weight_kg": 単重量(kg),
             "qty": 数量, "piece_cost": 仕入先見積りの枚単価（製品価格、円）}, ...] 最大8件まで。
        客先単価（円/kg）はテンプレート側の固定数式（原価単価+markup_per_kg）で自動計算される。
    markup_per_kg: 円/kgの定額乗せ（％上乗せではない）。P16セルに書き込む。
    freight_customer_price: 運賃の客先向け金額（円）。省略時は運賃行を空欄のままにする
        （通常は仕入運賃に対する掛け率で別途計算し、この引数に渡す）。
    """
    missing = []
    _require(customer_name, "客先名", missing)
    _require(contact_name, "担当者名", missing)
    _require(item_title, "件名（品名・物件名）", missing)
    if markup_per_kg is None:
        missing.append("円/kg乗せ（乗せ代）")
    if missing:
        raise MissingRequiredFieldError(
            f"チャットでの確認が必要な項目が未指定です: {', '.join(missing)}。"
        )

    if len(items) > len(ITEM_ROWS):
        raise TooManyItemsError(
            f"テンプレートの明細枠は{len(ITEM_ROWS)}件までです（{len(items)}件指定されました）。"
        )

    ws["B2"] = customer_name
    ws["B3"] = contact_name
    ws["D10"] = item_title
    ws["C11"] = f"受渡場所　及び条件　：{receiving_place}"
    ws["C12"] = f"決済条件               ：{payment_terms}"
    if note:
        ws["C14"] = f"備考：{note}"
    if quote_number:
        ws[QUOTE_NUMBER_CELL] = f"{QUOTE_NUMBER_LABEL}：{quote_number}"

    ws[MARKUP_PER_KG_CELL] = markup_per_kg

    for row, item in zip(ITEM_ROWS, items):
        ws[f"B{row}"] = item["thickness"]
        ws[f"C{row}"] = item["width"]
        ws[f"D{row}"] = item["length"]
        ws[f"E{row}"] = item["unit_weight_kg"]
        ws[f"F{row}"] = item["qty"]
        ws[f"O{row}"] = item["piece_cost"]

    if freight_customer_price is not None:
        ws["J26"] = freight_customer_price

    if stamp_path:
        _add_stamp(
            ws, stamp_path,
            cell=stamp_cell, offset_x_px=stamp_offset_x_px,
            offset_y_px=stamp_offset_y_px, size_px=stamp_size_px,
        )


def fill_steel_quote_template(template_path: Path, output_path: Path, **kwargs) -> Path:
    """テンプレートファイルを1件分の見積データで埋め、単独のxlsxとして保存する。
    kwargsはfill_steel_quote_sheet()と同じ。"""
    import shutil

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(template_path, output_path)

    wb = openpyxl.load_workbook(output_path)
    fill_steel_quote_sheet(wb.active, **kwargs)

    wb.save(output_path)
    return output_path
